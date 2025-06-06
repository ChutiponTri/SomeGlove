from openpyxl import load_workbook
from datetime import datetime
from threading import Lock
import sqlite3
import pytz

# data = {
#     "name": self.ชื่อ.text(),
#     "accessor": self.ผู้ดูแล.text(),
#     "gender": self.gender(),
#     "birth": self.birthday(),
#     "weight": float(self.weight_input.text()) if self.weight_input.text() != "" else "",
#     "height": float(self.height_input.text()) if self.height_input.text() != "" else "",
#     "ans1": self.ans1_check(),
#     "ans2": self.ans2_check(),
#     "ans3": self.ans3_check(),
#     "ans4": self.ans4_check(),
#     "ans5": self.ans5_check(),
#     "ans6": self.ans6_check(),
# }

lock = Lock()

class Database():
    def __init__(self):
        self.create_user_table()
        self.create_score_table()
        self.create_overview_table()
        self.create_data_table()
        self.create_condition_table()
        self.create_strength_table()
        self.create_process_table()
        self.create_feedback_table()

    # Function to Input User into Database
    def input_data(self, data:dict):
        #Insert User into Database
        self.insert_user(data)
        self.insert_condition(data)
        print("Create New User Successfully")

        # Load Workbook and Nescessary Adjust
        workbook = load_workbook("excel/template.xlsx")
        user = workbook["ข้อมูลทั่วไป"]
        user["B2"] = data["name"]
        user["B3"] = data["accessor"]
        user["B5"] = data["birth"]
        user["D5"] = int(datetime.now().year) - int(data["birth"].split("/")[-1])
        user["F5"] = data["gender"]
        user["B6"] = data["height"]
        user["D6"] = data["weight"]
        user["F6"] = (data["weight"]) / ((data["height"] / 100)**2)
        user["F8"] = data["ans1"]
        user["F9"] = data["ans2"]
        user["F10"] = data["ans3"]
        user["F11"] = data["ans4"]
        user["F12"] = data["ans5"]
        user["F13"] = data["ans6"]

        # Save Workbook
        workbook.save("excel/%s.xlsx" % data["name"])
        print("Created File %s.xlsx Successfully" % data["name"])

    # Function to Create User Table
    def create_user_table(self):
        conn = sqlite3.connect("database/glove.db")
        cursor = conn.cursor()
        query = """CREATE TABLE IF NOT EXISTS users(
            hn TEXT PRIMARY KEY NOT NULL,
            accessor TEXT NOT NULL,
            birth DATE NOT NULL,
            age INTEGER NOT NULL,
            gender TEXT NOT NULL,
            height REAL,
            weight REAL,
            bmi REAL,
            createdAt DATETIME DEFAULT CURRENT_TIMESTAMP
        )"""
        cursor.execute(query)
        conn.commit()
        conn.close()

    # Function to Create Score Table
    def create_score_table(self):
        conn = sqlite3.connect("database/glove.db")
        cursor = conn.cursor()
        query = """CREATE TABLE IF NOT EXISTS score(
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            hn TEXT NOT NULL,
            game TEXT NOT NULL,
            score INTEGER NOT NULL,
            start DATETIME NOT NULL,
            finish DATETIME NOT NULL,
            accuracy TEXT,
            note TEXT,
            FOREIGN KEY (hn) REFERENCES users(hn)
        )"""
        cursor.execute(query)
        conn.commit()
        conn.close()

    # Function to Create Overview Table
    def create_overview_table(self):
        conn = sqlite3.connect("database/glove.db")
        cursor = conn.cursor()
        query = """CREATE TABLE IF NOT EXISTS overview(
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            start DATETIME NOT NULL,
            finish DATETIME NOT NULL,
            hn TEXT NOT NULL,
            duration REAL,
            grip INTEGER,
            meanpressure REAL,
            maxpressure REAL,
            power REAL,
            calorie REAL,
            game TEXT,
            FOREIGN KEY (hn) REFERENCES users(hn)
        )"""
        cursor.execute(query)
        conn.commit()
        conn.close()

    # Function to Create Raw Data Table
    def create_data_table(self):
        conn = sqlite3.connect("database/glove.db")
        cursor = conn.cursor()
        query = """CREATE TABLE IF NOT EXISTS data(
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            hn TEXT NOT NULL,
            pressure REAL NOT NULL,
            createdAt DATETIME DEFAULT CURRENT_TIMESTAMP,
            FOREIGN KEY (hn) REFERENCES users(hn)
        )"""
        cursor.execute(query)
        conn.commit()
        conn.close()

    # Function to Create Condition Table
    def create_condition_table(self):
        conn = sqlite3.connect("database/glove.db")
        cursor = conn.cursor()
        query = """CREATE TABLE IF NOT EXISTS condition(
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            hn TEXT NOT NULL,
            ans1 TEXT NOT NULL,
            ans2 TEXT NOT NULL,
            ans3 TEXT NOT NULL,
            ans4 TEXT NOT NULL,
            ans5 TEXT NOT NULL,
            ans6 TEXT NOT NULL,
            createdAt DATETIME DEFAULT CURRENT_TIMESTAMP,
            FOREIGN KEY (hn) REFERENCES users(hn)
        )"""
        cursor.execute(query)
        conn.commit()
        conn.close()

    # Function to Create Strength Table
    def create_strength_table(self):
        conn = sqlite3.connect("database/glove.db")
        cursor = conn.cursor()
        query = """CREATE TABLE IF NOT EXISTS strength (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            hn TEXT NOT NULL UNIQUE,
            beforeOne REAL,
            afterOne REAL,
            beforeTwo REAL,
            afterTwo REAL,
            beforeBest REAL,
            afterBest REAL,
            createdAt DATETIME DEFAULT CURRENT_TIMESTAMP,
            updatedAt DATETIME DEFAULT CURRENT_TIMESTAMP,
            FOREIGN KEY (hn) REFERENCES users(hn)
        )"""
        cursor.execute(query)
        cursor.executescript("""
            DROP TRIGGER IF EXISTS strength_updatedAt;
            CREATE TRIGGER strength_updatedAt
            AFTER UPDATE ON strength
            FOR EACH ROW
            BEGIN
                UPDATE strength
                SET updatedAt = CURRENT_TIMESTAMP
                WHERE id = NEW.id;
            END;
        """)
        conn.commit()
        conn.close()

    # Function to Create Process Table
    def create_process_table(self):
        conn = sqlite3.connect("database/glove.db")
        cursor = conn.cursor()
        query = """CREATE TABLE IF NOT EXISTS process (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            hn TEXT NOT NULL UNIQUE,
            oneBeforeLeft INTEGER,
            oneBeforeRight INTEGER,
            oneAfterLeft INTEGER,
            oneAfterRight INTEGER,
            twoBeforeLeft INTEGER,
            twoBeforeRight INTEGER,
            twoAfterLeft INTEGER,
            twoAfterRight INTEGER,
            threeBeforeLeft INTEGER,
            threeBeforeRight INTEGER,
            threeAfterLeft INTEGER,
            threeAfterRight INTEGER,
            fourBeforeLeft INTEGER,
            fourBeforeRight INTEGER,
            fourAfterLeft INTEGER,
            fourAfterRight INTEGER,
            fiveBeforeLeft INTEGER,
            fiveBeforeRight INTEGER,
            fiveAfterLeft INTEGER,
            fiveAfterRight INTEGER,
            sixBeforeLeft INTEGER,
            sixBeforeRight INTEGER,
            sixAfterLeft INTEGER,
            sixAfterRight INTEGER,
            sevenBeforeLeft INTEGER,
            sevenBeforeRight INTEGER,
            sevenAfterLeft INTEGER,
            sevenAfterRight INTEGER,
            createdAt DATETIME DEFAULT CURRENT_TIMESTAMP,
            updatedAt DATETIME DEFAULT CURRENT_TIMESTAMP,
            FOREIGN KEY (hn) REFERENCES users(hn)
        )"""
        cursor.execute(query)
        cursor.executescript("""
            DROP TRIGGER IF EXISTS trg_updatedAt;
            CREATE TRIGGER trg_updatedAt
            AFTER UPDATE ON process
            FOR EACH ROW
            BEGIN
                UPDATE process
                SET updatedAt = CURRENT_TIMESTAMP
                WHERE id = NEW.id;
            END;
        """)
        conn.commit()
        conn.close()

    # Function to Create Feedback Table
    def create_feedback_table(self):
        conn = sqlite3.connect("database/glove.db")
        cursor = conn.cursor()
        query = """CREATE TABLE IF NOT EXISTS feedback (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            hn TEXT NOT NULL UNIQUE,
            one INTEGER NOT NULL,
            two INTEGER NOT NULL,
            three INTEGER NOT NULL,
            four INTEGER NOT NULL,
            five INTEGER NOT NULL,
            six INTEGER NOT NULL,
            seven INTEGER NOT NULL,
            eight INTEGER NOT NULL,
            nine INTEGER NOT NULL,
            ten INTEGER NOT NULL,
            totalOne INTEGER NOT NULL,
            totalTwo INTEGER NOT NULL,
            totalThree INTEGER NOT NULL,
            totalFour INTEGER NOT NULL,
            totalFive INTEGER NOT NULL,
            createdAt DATETIME DEFAULT CURRENT_TIMESTAMP,
            updatedAt DATETIME DEFAULT CURRENT_TIMESTAMP,
            FOREIGN KEY (hn) REFERENCES users(hn)
        )"""
        cursor.execute(query)
        cursor.executescript("""
            DROP TRIGGER IF EXISTS feedback_updatedAt;
            CREATE TRIGGER feedback_updatedAt
            AFTER UPDATE ON feedback
            FOR EACH ROW
            BEGIN
                UPDATE feedback
                SET updatedAt = CURRENT_TIMESTAMP
                WHERE id = NEW.id;
            END;
        """)
        conn.commit()
        conn.close()

    # Function to Get All Users
    @staticmethod
    def get_all_usernames():
        conn = sqlite3.connect("database/glove.db")
        cursor = conn.cursor()
        query = "SELECT hn FROM users"
        cursor.execute(query)
        users = cursor.fetchall()
        conn.commit()
        conn.close()
        return [user[0] for user in users]

    # Function to Check if User Exists
    def get_exist_user(self, username):
        conn = sqlite3.connect("database/glove.db")
        cursor = conn.cursor()
        query = "SELECT * FROM users WHERE hn = ?"
        cursor.execute(query, (username,))
        user = cursor.fetchone()
        conn.commit()
        conn.close()
        return user

    # Function to Insert into Users Table
    def insert_user(self, data: dict):
        conn = sqlite3.connect("database/glove.db")
        cursor = conn.cursor()
        age = int(datetime.now().year) - int(data["birth"].split("/")[-1])
        bmi = (data["weight"]) / ((data["height"] / 100)**2)
        query = "INSERT INTO users (hn, accessor, birth, age, gender, height, weight, bmi) VALUES (?, ?, ?, ?, ?, ?, ?, ?)"
        cursor.execute(query, (data["name"], data["accessor"], data["birth"], age, data["gender"], data["height"], data["weight"], bmi))
        conn.commit()
        conn.close()

    # Function to Insert User Condition
    def insert_condition(self, data: dict):
        conn = sqlite3.connect("database/glove.db")
        cursor = conn.cursor()
        query = "INSERT INTO condition (hn, ans1, ans2, ans3, ans4, ans5, ans6) VALUES (?, ?, ?, ?, ?, ?, ?)"
        cursor.execute(query, (data["name"], data["ans1"], data["ans2"], data["ans3"], data["ans4"], data["ans5"], data["ans6"]))
        conn.commit()
        conn.close()

    # Function to Create Login Table
    @staticmethod
    def create_current_user():
        conn = sqlite3.connect("database/login.db")
        c = conn.cursor()
        query = """CREATE TABLE IF NOT EXISTS current (
            Username TEXT,
            Timestamp TIMESTAMP DEFAULT CURRENT_TIMESTAMP,
            Status TEXT
        )"""
        c.execute(query)
        conn.commit()
        conn.close()

    # Function to Update Login Table
    @staticmethod
    def current_user(username):   
        conn = sqlite3.connect("database/login.db")
        c = conn.cursor()
        timestamp = datetime.now().strftime("%Y-%m-%d %H-%M-%S")
        status = "login"
        c.execute("INSERT INTO current VALUES(?, ?, ?)", (username, timestamp, status))
        conn.commit()
        conn.close()

    # Function to Check Login Table
    @staticmethod
    def get_current_user():
        conn = sqlite3.connect("database/login.db")
        c = conn.cursor()
        status = "login"
        c.execute("SELECT Username FROM current WHERE Status=? ORDER BY rowid DESC LIMIT 1", (status,))
        user = c.fetchone()
        conn.commit()
        conn.close()
        return user[0]

    # Function to Del Login Table
    @staticmethod
    def del_current_user(username):
        conn = sqlite3.connect("database/login.db")
        c = conn.cursor()
        timestamp = datetime.now().strftime("%Y-%m-%d %H-%M-%S")
        status = "logout"
        c.execute("INSERT INTO current VALUES(?,?,?)", (username, timestamp, status))
        print(f"Logout from {username}")
        conn.commit()
        conn.close()

    # Function to Insert Score to Database
    @staticmethod
    def insert_score(username, game, score, start, finish, accuracy, note: str):
        conn = sqlite3.connect("database/glove.db")
        cursor = conn.cursor()
        query = "INSERT INTO score (hn, game, score, start, finish, accuracy, note) VALUES (?, ?, ?, ?, ?, ?, ?)"
        cursor.execute(query, (username, game, score, start, finish, accuracy, note))
        conn.commit()
        conn.close()

    # Function to Insert Overview
    @staticmethod
    def insert_overview(start, finish, username, duration, grip, meanpressure, maxpressure, power, calorie, game):
        conn = sqlite3.connect("database/glove.db")
        cursor = conn.cursor()
        query = "INSERT INTO overview (start, finish, hn, duration, grip, meanpressure, maxpressure, power, calorie, game) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?)"
        cursor.execute(query, (start, finish, username, duration, grip, meanpressure, maxpressure, power, calorie, game))
        conn.commit()
        conn.close()  

    # Function to Insert Data
    @staticmethod
    def insert_data(username, pressure):
        conn = sqlite3.connect("database/glove.db")
        cursor = conn.cursor()
        query = "INSERT INTO data (hn, pressure) VALUES (?, ?)"
        cursor.execute(query, (username, pressure))
        conn.commit()
        conn.close()  

    # Function to Insert Strength 
    @staticmethod
    def insert_strength(username, **kwargs):
        valid_keys = [
            "beforeOne", "afterOne",
            "beforeTwo", "afterTwo",
            "beforeBest", "afterBest"
        ]
        data = {k: v for k, v in kwargs.items() if k in valid_keys}
        if not data:   
            print("⚠️ No valid strength fields provided.")
            return
        
        data["hn"] = username
        keys = ", ".join(data.keys())
        placeholders = ", ".join("?" for _ in data)
        update_clause = ", ".join(f"{k}=excluded.{k}" for k in data if k != "hn")

        query = f"""
            INSERT INTO strength ({keys})
            VALUES ({placeholders})
            ON CONFLICT(hn) DO UPDATE SET {update_clause};
        """

        conn = sqlite3.connect("database/glove.db")
        conn.execute("PRAGMA foreign_keys = ON")
        cursor = conn.cursor()
        cursor.execute(query, list(data.values()))
        conn.commit()
        conn.close()

    # Function to Insert Process
    @staticmethod
    def insert_process(username, **kwargs):
        valid_keys = [
            "oneBeforeLeft", "oneBeforeRight", "oneAfterLeft", "oneAfterRight",
            "twoBeforeLeft", "twoBeforeRight", "twoAfterLeft", "twoAfterRight",
            "threeBeforeLeft", "threeBeforeRight", "threeAfterLeft", "threeAfterRight",
            "fourBeforeLeft", "fourBeforeRight", "fourAfterLeft", "fourAfterRight",
            "fiveBeforeLeft", "fiveBeforeRight", "fiveAfterLeft", "fiveAfterRight",
            "sixBeforeLeft", "sixBeforeRight", "sixAfterLeft", "sixAfterRight",
            "sevenBeforeLeft", "sevenBeforeRight", "sevenAfterLeft", "sevenAfterRight"
        ]
        data = {k: v for k, v in kwargs.items() if k in valid_keys}
        if not data:
            print("No valid process fields provided.")
            return
        
        data["hn"] = username  
        keys = ", ".join(data.keys())
        placeholders = ", ".join(["?" for _ in data])
        update_clause = ", ".join([f"{key}=excluded.{key}" for key in data if key != "hn"])

        query = f"""
            INSERT INTO process ({keys})
            VALUES ({placeholders})
            ON CONFLICT(hn) DO UPDATE SET {update_clause};
        """
        conn = sqlite3.connect("database/glove.db")
        conn.execute("PRAGMA foreign_keys = ON")
        cursor = conn.cursor()
        cursor.execute(query, list(data.values()))
        conn.commit()
        conn.close()

    # Function to Insert Feedback
    @staticmethod
    def insert_feedback(username, **kwargs):
        valid_keys = [
            "one", "two", "three", "four", "five", "six", "seven", "eight", "nine", "ten",
            "totalOne", "totalTwo", "totalThree", "totalFour", "totalFive"
        ]
        data = {k: v for k, v in kwargs.items() if k in valid_keys}
        if len(data) != 15:
            print("⚠️ Must provide all 15 feedback fields.")
            return
        
        data["hn"] = username
        keys = ", ".join(data.keys())
        placeholders = ", ".join("?" for _ in data)
        update_clause = ", ".join(f"{k}=excluded.{k}" for k in data if k != "hn")

        query = f"""
            INSERT INTO feedback ({keys})
            VALUES ({placeholders})
            ON CONFLICT(hn) DO UPDATE SET {update_clause};
        """

        conn = sqlite3.connect("database/glove.db")
        conn.execute("PRAGMA foreign_keys = ON")
        cursor = conn.cursor()
        cursor.execute(query, list(data.values()))
        conn.commit()
        conn.close()

    # Function to Get Strength 
    @staticmethod
    def get_strength(username):
        conn = sqlite3.connect("database/glove.db")
        cursor = conn.cursor()
        query = "SELECT * FROM strength WHERE hn = ?"
        cursor.execute(query, (username,))
        data = cursor.fetchone()
        conn.commit()
        conn.close()
        return data
    
    # Function to Get Process
    @staticmethod
    def get_process(username):
        conn = sqlite3.connect("database/glove.db")
        cursor = conn.cursor()
        query = "SELECT * FROM process WHERE hn = ?"
        cursor.execute(query, (username,))
        data = cursor.fetchone()
        conn.commit()
        conn.close()
        return data
    
    # Function to Get Feedback
    @staticmethod
    def get_feedback(username):
        conn = sqlite3.connect("database/glove.db")
        cursor = conn.cursor()
        query = "SELECT * FROM feedback WHERE hn = ?"
        cursor.execute(query, (username,))
        data = cursor.fetchone()
        conn.commit()
        conn.close()
        return data

    # Function to Get Overview Start Time
    @staticmethod
    def get_start_index(username):
        conn = sqlite3.connect("database/glove.db")
        cursor = conn.cursor()
        query = "SELECT start, finish FROM overview WHERE hn = ?"
        cursor.execute(query, (username,))
        data = cursor.fetchall()
        conn.commit()
        conn.close()
        return {start: finish for start, finish in data}
    
    # Function to Get Pressure Data
    @staticmethod
    def get_pressure(username, start: str, finish: str):
        conn = sqlite3.connect("database/glove.db")
        cursor = conn.cursor()
        local = pytz.timezone("Asia/Bangkok")
        start_local = local.localize(datetime.strptime(start, "%Y-%m-%d %H:%M:%S"))
        start_utc = start_local.astimezone(pytz.utc).strftime("%Y-%m-%d %H:%M:%S")
        finish_local = local.localize(datetime.strptime(finish, "%Y-%m-%d %H:%M:%S"))
        finish_utc = finish_local.astimezone(pytz.utc).strftime("%Y-%m-%d %H:%M:%S")
        query = "SELECT createdAt, pressure FROM data WHERE hn = ? AND createdAt >= ? AND createdAt <= ?"
        cursor.execute(query, (username, start_utc, finish_utc))
        data = cursor.fetchall()
        conn.commit()
        conn.close()
        timestamps = [i for [i, _] in data]
        pressures  = [j for [_, j] in data]
        return {"timestamp": timestamps, "pressure": pressures}
    
    # Function to Get Overview Data
    @staticmethod
    def get_overview(username):
        conn = sqlite3.connect("database/glove.db")
        cursor = conn.cursor()
        query = "SELECT * FROM overview WHERE hn = ?"
        cursor.execute(query, (username,))
        data = cursor.fetchall()
        conn.commit()
        conn.close()
        
        return {
            "เวลาขณะเริ่มต้น": [i[1] for i in data],
            "เวลาขณะเสร็จสิ้น": [i[2] for i in data],
            "เวลาที่ใช้ (วินาที)" : [i[4] for i in data],
            "จำนวนครั้งที่กำมือ" : [i[5] for i in data],
            "ความดันเฉลี่ย (kPa)" : [i[6] for i in data],
            "ความดันสูงสุด (kPa)" : [i[7] for i in data],
            "พลังงาน (W)": [i[8] for i in data],
            "แคลอรี่ (kcal)" : [i[9] for i in data],
            "เกมที่ทดสอบ" : [i[10] for i in data],
        }
    
    # Function to Get Game Score
    @staticmethod
    def get_game(username):
        conn = sqlite3.connect("database/glove.db")
        cursor = conn.cursor()
        query = "SELECT * FROM score WHERE hn = ?"
        cursor.execute(query, (username,))
        data = cursor.fetchall()
        conn.commit()
        conn.close()

        return {
            "เกม": [i[2] for i in data],
            "คะแนน": [i[3] for i in data],
            "เวลาเริ่มต้น": [i[4] for i in data],
            "เวลาสิ้นสุด": [i[5] for i in data],
            "ความแม่นยำ": [i[6] for i in data],
            "หมายเหตุ": [i[7] for i in data]
        }
    
if __name__ == "__main__":
    Database.create_current_user()
    # Database.current_user("Ton")
    # Database.del_current_user("Ton")
    name = Database.get_current_user()
    print(name)
